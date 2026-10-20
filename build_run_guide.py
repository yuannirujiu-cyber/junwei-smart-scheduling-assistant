from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = "君威智能排产助手运行说明.docx"


def set_font(run, size=None, bold=None, color=None):
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    props.append(shd)


def borders(cell, color="D9D9D9"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def cell_text(cell, text, bold=False, color=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run(text)
    set_font(r, 9.5, bold, color)
    borders(cell)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(15 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_font(r, 15 if level == 1 else 12, True, (0, 0, 0))
    return p


def add_body(doc, text, indent=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.28
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.25)
    r = p.add_run(text)
    set_font(r, 10.5, False, (35, 43, 58))
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_font(r, 10.5, False, (35, 43, 58))
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run(text)
    set_font(r, 10.5, False, (35, 43, 58))
    return p


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(30, 64, 175)
    return p


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)

    styles = doc.styles
    styles["Normal"].font.name = "Microsoft YaHei"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("君威智能排产助手运行说明")
    set_font(r, 22, True, (0, 0, 0))
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    sr = subtitle.add_run("适用于本地网页部署与日常排产操作")
    set_font(sr, 11, False, (89, 102, 125))
    add_body(doc, "本说明用于帮助操作人员启动君威智能排产助手，并完成订单导入、产品定额更换、排产、快速重排及导出等日常工作。", True)

    add_heading(doc, "1 运行环境")
    add_bullet(doc, "已安装 Node.js 18 或更高版本。")
    add_bullet(doc, "本项目不需要安装 npm 依赖，也不需要联网。")
    add_bullet(doc, "请通过 HTTP 服务访问网页，不能直接双击打开 index.html。")

    add_heading(doc, "2 启动网页")
    add_body(doc, "在项目根目录执行以下命令：")
    add_code(doc, "node tools/serve.mjs")
    add_body(doc, "浏览器访问：")
    add_code(doc, "http://127.0.0.1:5178/web/index.html")
    add_body(doc, "如果 5178 端口已被占用，可换用其他端口：")
    add_code(doc, "$env:PORT=5180; node tools/serve.mjs")
    add_body(doc, "随后访问 http://127.0.0.1:5180/web/index.html。")

    add_heading(doc, "3 首次使用")
    add_body(doc, "网页首次打开时，内置产品定额数据已经加载完成，排产参考日期默认为 2026/08/03。")
    for item in [
        "进入排产工作台。",
        "点击上传订单。",
        "选择订单 xlsx 或 csv 文件，单个文件最大 20 MB。",
        "在字段映射窗口确认订单号、规格、业务数量、预发货日等必填字段。",
        "确认导入后，系统会自动生成首排计划。",
    ]:
        add_number(doc, item)
    add_body(doc, "订单导入后，可在甘特图顶部输入完整订单号或设备编号进行精确筛选。当输入内容是设备编号时，设备匹配优先，不会混入同号订单。")

    add_heading(doc, "4 更换产品定额")
    add_body(doc, "如需使用新的产品定额，可按以下步骤操作：")
    for item in [
        "点击页面顶部更换产品定额。",
        "确认替换后，选择产品定额 xlsx 文件。",
        "确认文件包含拉丝、捻股、合绳三类工作表。",
        "在预览中核对设备数量及产能矩阵行数，确认后完成替换。",
        "点击重新生成，得到基于新定额的计划。",
    ]:
        add_number(doc, item)
    add_body(doc, "已保存的计划会保留并标识为基于旧数据。矩阵中的空值或非正产能会被视为该设备不适配该规格。")

    add_heading(doc, "5 日常排产操作")
    operations = [
        ("日期", "修改排产参考日期只影响甘特图的参考线与定位，不改变排产结果。"),
        ("重新生成", "按当前订单、产品定额和设置重新计算计划。"),
        ("甘特图", "可在设备视角和订单视角间切换；全部、拉丝、捻股、合绳按钮用于按工序显示设备数据。"),
        ("订单中心", "查看订单清单、排产状态以及风险与异常；两侧列表可独立滚动。"),
        ("材料库", "维护材料库存和有无状态；缺料订单会进入阻断风险。"),
        ("清空订单信息", "清空当前订单和计划，但保留当前产品定额数据。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.35)
    table.columns[1].width = Inches(5.65)
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ("功能", "说明")):
        shade(cell, "1F4E78")
        cell_text(cell, text, True, (255, 255, 255), WD_ALIGN_PARAGRAPH.CENTER)
    for i, row in enumerate(operations):
        cells = table.add_row().cells
        if i % 2 == 1:
            for cell in cells:
                shade(cell, "F4F8FC")
        cell_text(cells[0], row[0], True, (35, 43, 58), WD_ALIGN_PARAGRAPH.CENTER)
        cell_text(cells[1], row[1])

    add_heading(doc, "6 异常快速重排")
    add_body(doc, "进入异常重排后，可处理紧急插单、设备故障和物料延迟。系统会以当前计划为基线，不直接覆盖原计划。")
    for item in [
        "先只释放直接受影响的未开工任务。",
        "如无法得到可行结果，再释放同设备后续冲突任务。",
        "仍无解时，继续释放关联订单和工序依赖任务。",
        "生成候选后查看变更订单、变更工序、设备变更、交期变化及影响范围扩展记录。",
        "点击采用此方案后，候选计划才会成为当前草稿；取消不会修改基线计划。",
    ]:
        add_number(doc, item)
    add_body(doc, "人工锁定任务和执行冻结任务不会被快速重排移动。")

    add_heading(doc, "7 风险与指标说明")
    add_bullet(doc, "风险与异常显示全部校验结果，包括工艺映射缺失、材料缺失、订单字段缺失和交期风险。")
    add_bullet(doc, "预计延期订单和硬交期风险只统计已能够完成排产并计算完工时间的订单。")
    add_bullet(doc, "无法排产的订单可能出现在风险面板中，但不会计入延期或硬交期指标；应优先处理风险面板中的阻断项。")

    add_heading(doc, "8 导出")
    add_body(doc, "点击顶部导出，在弹窗中选择所需文件类型并再次点击导出。可导出排产清单、工单、甘特图、风险清单、当前筛选 CSV、负荷分析报告和审计日志等文件。")

    add_heading(doc, "9 常见问题")
    issues = [
        ("页面无法打开或一直加载", "确认使用 node tools/serve.mjs 启动，并通过 http://127.0.0.1 访问。"),
        ("导入订单后无法排产", "检查字段映射是否包含订单号、规格、业务数量和预发货日；再查看风险面板中的阻断项。"),
        ("产品定额导入失败", "确认文件为 xlsx、小于 20 MB，且包含拉丝、捻股、合绳相关工作表。"),
        ("甘特图没有显示预期订单或设备", "输入完整订单号或设备编号；点击全部清除工序筛选。"),
        ("更换定额后计划仍是旧结果", "点击重新生成，旧版本不会自动被覆盖。"),
        ("需要恢复原计划", "在异常重排页面不要采用候选方案；原基线版本始终保留。"),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(4.8)
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ("现象", "处理方式")):
        shade(cell, "1F4E78")
        cell_text(cell, text, True, (255, 255, 255), WD_ALIGN_PARAGRAPH.CENTER)
    for i, row in enumerate(issues):
        cells = table.add_row().cells
        if i % 2 == 1:
            for cell in cells:
                shade(cell, "F4F8FC")
        cell_text(cells[0], row[0], True)
        cell_text(cells[1], row[1])

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("君威智能排产助手")
    set_font(fr, 8.5, False, (118, 130, 150))
    doc.save(OUT)


if __name__ == "__main__":
    main()
